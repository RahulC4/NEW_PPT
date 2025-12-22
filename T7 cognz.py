# =============================================
# generate_ppt_cognizant.py  (FINAL – SAFE)
# =============================================
import os
import uuid
from pptx import Presentation
from pptx.util import Pt
from utils import logger


COGNIZANT_TEMPLATE_PATH = "templates/Cognizant.pptx"


# ------------------------------------------------------------
# Helpers
# ------------------------------------------------------------
def clear_text_frame(tf):
    tf.clear()
    tf.word_wrap = True


def set_title(slide, text):
    if slide.shapes.title:
        slide.shapes.title.text = text


def set_bullets(slide, bullets):
    if len(slide.placeholders) < 2:
        return

    tf = slide.placeholders[1].text_frame
    clear_text_frame(tf)

    for i, bullet in enumerate(bullets):
        p = tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(18)


# ------------------------------------------------------------
# MAIN GENERATOR
# ------------------------------------------------------------
def generate_presentation(payload):
    """
    Generates PPT using Cognizant template
    Replaces content ONLY (no slide cloning)
    """

    preview_slides = payload.get("preview_slides")
    if not preview_slides:
        raise ValueError("No preview slides found")

    if not os.path.exists(COGNIZANT_TEMPLATE_PATH):
        raise FileNotFoundError("Cognizant template not found")

    prs = Presentation(COGNIZANT_TEMPLATE_PATH)

    # ------------------------------------------------------------
    # ASSUMPTIONS ABOUT TEMPLATE
    # ------------------------------------------------------------
    # slide 0 -> title slide
    # slide 1 -> content template
    # slide -1 -> thank you (kept as-is)
    # ------------------------------------------------------------

    content_template_index = 1
    thank_you_index = len(prs.slides) - 1

    # ------------------------------------------------------------
    # TITLE SLIDE
    # ------------------------------------------------------------
    title_slide = prs.slides[0]
    set_title(title_slide, preview_slides[0]["title"])

    # ------------------------------------------------------------
    # CONTENT SLIDES
    # ------------------------------------------------------------
    for idx, preview in enumerate(preview_slides[1:], start=1):

        # Stop before Thank You
        if idx >= thank_you_index:
            break

        slide = prs.slides[idx]  # ✅ SINGLE SLIDE OBJECT

        title = preview.get("title", "").strip()
        bullets = preview.get("bullets", [])

        # Normalize bullets (CRITICAL)
        bullets = [
            b.strip()
            for b in bullets
            if isinstance(b, str) and b.strip()
        ]

        set_title(slide, title)
        set_bullets(slide, bullets)

    # ------------------------------------------------------------
    # SAVE
    # ------------------------------------------------------------
    os.makedirs("generated", exist_ok=True)
    out_path = f"generated/cognizant_{uuid.uuid4().hex[:6]}.pptx"
    prs.save(out_path)

    return out_path
