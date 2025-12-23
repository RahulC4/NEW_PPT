# ============================================================
# generate_ppt_cognizant.py
# Cognizant Template | Preview-driven | TEXT-ONLY (No Images)
# ============================================================

import os
import uuid
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from utils import logger


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
COGNIZANT_TEMPLATE = os.path.join(TEMPLATES_DIR, "Cognizant.pptx")


# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def delete_slide(prs, index):
    slide_id = prs.slides._sldIdLst[index].rId
    prs.part.drop_rel(slide_id)
    del prs.slides._sldIdLst[index]


def remove_empty_placeholders(slide):
    """
    Removes 'Click to add...' placeholders AFTER content is filled
    """
    for shape in list(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip().lower()
            if txt.startswith("click to add"):
                slide.shapes._spTree.remove(shape._element)


def update_month_year(slide):
    """
    Updates date text to current Month YYYY
    """
    current = datetime.now().strftime("%B %Y")
    for shape in slide.shapes:
        if shape.has_text_frame:
            if any(y in shape.text for y in ["2023", "2024", "2025", "2026"]):
                shape.text = current


def set_title_slide_title(prs, slide, text):
    """
    FIRST SLIDE ONLY
    Full-width, white title, single line
    """
    for shape in list(slide.shapes):
        if shape.is_placeholder:
            slide.shapes._spTree.remove(shape._element)

    tb = slide.shapes.add_textbox(
        left=Inches(0.5),
        top=Inches(2.6),
        width=prs.slide_width - Inches(1),
        height=Pt(120),
    )

    tf = tb.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)


def apply_footer(slide):
    """
    Apply footer text to template footer placeholder
    """
    year = datetime.now().year
    for shape in slide.shapes:
        if shape.has_text_frame:
            if shape.text.strip().lower() == "footer":
                shape.text = f"© {year} Cognizant"


# ------------------------------------------------------------
# MAIN GENERATOR
# ------------------------------------------------------------
def generate_presentation_cognizant(payload):
    """
    Cognizant PPT generation
    - Preview driven
    - Title & Thank You preserved
    - Content slides fixed (bullets + footer)
    """

    slides = payload.get("slides")
    if not slides:
        raise ValueError("No preview slides found")

    prs = Presentation(COGNIZANT_TEMPLATE)

    # Capture layouts BEFORE deleting
    title_layout = prs.slides[0].slide_layout
    content_layout = prs.slides[1].slide_layout
    thankyou_layout = prs.slides[-1].slide_layout

    # Remove all template slides
    for i in reversed(range(len(prs.slides))):
        delete_slide(prs, i)

    # --------------------------------------------------------
    # 1️⃣ TITLE SLIDE
    # --------------------------------------------------------
    title_slide = prs.slides.add_slide(title_layout)
    set_title_slide_title(prs, title_slide, slides[0].get("title", ""))
    update_month_year(title_slide)
    apply_footer(title_slide)
    remove_empty_placeholders(title_slide)

    # --------------------------------------------------------
    # 2️⃣ CONTENT SLIDES  ✅ FULL FIX
    # --------------------------------------------------------
    for slide_data in slides[1:]:
        slide = prs.slides.add_slide(content_layout)

        title = slide_data.get("title", "")
        bullets = [
            b.strip()
            for b in (slide_data.get("bullets") or [])
            if isinstance(b, str) and b.strip()
        ]

        # ---- TITLE (force white, DO NOT recreate) ----
        if slide.shapes.title:
            tf = slide.shapes.title.text_frame
            p = tf.paragraphs[0]
            p.text = title
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)

        # ---- BODY PLACEHOLDER (idx = 1 is KEY) ----
        body = None
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 1:
                body = ph
                break

        if body and body.has_text_frame:
            tf = body.text_frame
            tf.clear()

            for i, bullet in enumerate(bullets):
                if i == 0:
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()

                p.text = bullet
                p.level = 0           # 🔥 fixes first bullet
                p.font.size = Pt(20)

        # ---- FOOTER ----
        apply_footer(slide)

        # ---- CLEANUP ----
        remove_empty_placeholders(slide)

    # --------------------------------------------------------
    # 3️⃣ THANK YOU SLIDE
    # --------------------------------------------------------
    thank_slide = prs.slides.add_slide(thankyou_layout)

    for shape in thank_slide.shapes:
        if shape.has_text_frame:
            shape.text = "Thank You"
            break

    apply_footer(thank_slide)
    remove_empty_placeholders(thank_slide)

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------
    os.makedirs("generated", exist_ok=True)
    out_path = f"generated/cognizant_{uuid.uuid4().hex[:6]}.pptx"
    prs.save(out_path)
    return out_path
