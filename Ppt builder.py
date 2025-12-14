# ppt/ppt_builder.py

import os
import tempfile
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from utils import logger
from llm.llm_utils import generate_image


# ============================================================
#  HELPER: Apply text to placeholder shape using metadata
# ============================================================
def fill_text_placeholder(shape, text: str, font_meta: dict):
    """
    Fill a shape's text frame while preserving template formatting.
    """
    try:
        tf = shape.text_frame
        tf.clear()

        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = text

        # Apply font styling if provided
        if font_meta:
            if font_meta.get("name"):
                run.font.name = font_meta["name"]
            if font_meta.get("size"):
                run.font.size = Pt(font_meta["size"])
            if font_meta.get("bold") is not None:
                run.font.bold = font_meta["bold"]
            if font_meta.get("italic") is not None:
                run.font.italic = font_meta["italic"]
            if font_meta.get("color"):
                try:
                    run.font.color.rgb = font_meta["color"]
                except:
                    pass

    except Exception as e:
        logger.error(f"Failed to set text placeholder: {e}")


# ============================================================
#  HELPER: Insert image into bounding box
# ============================================================
def insert_image(slide, img_bytes, ph_meta):
    """
    Insert AI-generated image into the correct placeholder bounding box.
    """
    if img_bytes is None:
        return

    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmp:
            tmp.write(img_bytes)
            tmp_path = tmp.name

        slide.shapes.add_picture(
            tmp_path,
            ph_meta["left"],
            ph_meta["top"],
            width=ph_meta["width"],
            height=ph_meta["height"]
        )

        os.remove(tmp_path)

    except Exception as e:
        logger.error(f"Image insertion failed: {e}")


# ============================================================
#  MAIN: Build PPT using templates + AI content
# ============================================================
def build_presentation(mapped_templates: list, content_profile: dict, output_path: str):
    """
    mapped_templates: Output from template_mapper.map_plan_to_templates
    content_profile:  Final user-filled JSON from qna_engine
    output_path:      Where to save PPT
    """

    prs = Presentation()  # blank deck (we will overwrite slide layouts)

    for item in mapped_templates:
        requested_slide = item["requested_slide"]
        template = item["template"]

        slide_type = requested_slide.get("slide_type")
        slide_content_fields = requested_slide.get("required_fields", [])

        # Fetch content from user + AI (already processed outside)
        final_text = ""
        for f in slide_content_fields:
            val = content_profile["fields"].get(f, "")
            if val:
                final_text += f"- {val}\n"

        # Create slide using blank layout then position shapes manually
        slide_layout = prs.slide_layouts[6]  # EMPTY_LAYOUT
        slide = prs.slides.add_slide(slide_layout)

        layout_meta = template.get("layout")
        placeholders = template.get("placeholders", [])

        if not placeholders:
            # Fallback — just add a text box
            logger.warning(f"No layout metadata for slide_type={slide_type}, using fallback.")
            left = top = Pt(50)
            width = Pt(900)
            height = Pt(500)
            box = slide.shapes.add_textbox(left, top, width, height)
            box.text = final_text
            continue

        # -------------------------------------------------------
        # Rebuild slide using metadata
        # -------------------------------------------------------
        for ph in placeholders:
            ph_left = ph["left"]
            ph_top = ph["top"]
            ph_w = ph["width"]
            ph_h = ph["height"]

            # TEXT placeholder
            if ph.get("has_text_frame"):
                textbox = slide.shapes.add_textbox(ph_left, ph_top, ph_w, ph_h)
                font_meta = ph.get("font", {})
                fill_text_placeholder(textbox, final_text, font_meta)

            # IMAGE placeholder (placeholder_type  = picture / picture-like)
            elif str(ph.get("placeholder_type")).lower().find("picture") != -1:
                # Generate image
                img_prompt = f"Professional illustration for slide type: {slide_type}"
                img_bytes = generate_image(img_prompt, size="1024x1024")
                insert_image(slide, img_bytes, ph)

            # Other shapes (rectangles, icons, tables)
            else:
                # We can extend this later for tables, shapes, charts...
                pass

    prs.save(output_path)
    logger.info(f"Final PPT generated at: {output_path}")
