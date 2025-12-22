# =============================================
# generate_ppt_cognizant.py
# =============================================
import os
import uuid
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import PP_PLACEHOLDER
from utils import logger

COGNIZANT_TEMPLATE = "templates/Cognizant.pptx"


# --------------------------------------------------
# Helpers
# --------------------------------------------------
def _get_title_tf(slide):
    for ph in slide.placeholders:
        if (
            ph.placeholder_format.type == PP_PLACEHOLDER.TITLE
            and hasattr(ph, "text_frame")
        ):
            return ph.text_frame
    return None


def _get_body_tf(slide):
    for ph in slide.placeholders:
        if (
            ph.placeholder_format.type == PP_PLACEHOLDER.BODY
            and hasattr(ph, "text_frame")
        ):
            return ph.text_frame
    return None


def _clear_all_text(slide):
    for shape in slide.shapes:
        if hasattr(shape, "text_frame"):
            shape.text_frame.clear()


# --------------------------------------------------
# MAIN GENERATOR
# --------------------------------------------------
def generate_presentation(payload):
    preview_slides = payload.get("slides")  # 🔥 preview slides ONLY

    if not preview_slides:
        raise ValueError("No preview slides found")

    prs = Presentation(COGNIZANT_TEMPLATE)

    # --------------------------------------------
    # Remove ALL sample content slides
    # Keep only the last THANK YOU slide
    # --------------------------------------------
    while len(prs.slides) > 1:
        prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])

    thank_you_slide = prs.slides[0]

    # --------------------------------------------
    # Generate slides from preview
    # --------------------------------------------
    for idx, slide_data in enumerate(preview_slides):
        title = slide_data.get("title", "").strip()
        bullets = [
            b.strip()
            for b in slide_data.get("bullets", [])
            if isinstance(b, str) and b.strip()
        ]

        slide = prs.slides.add_slide(prs.slide_layouts[1])

        _clear_all_text(slide)

        # ---- Title ----
        title_tf = _get_title_tf(slide)
        if title_tf:
            title_tf.text = title

        # ---- Bullets ----
        body_tf = _get_body_tf(slide)
        if body_tf and bullets:
            body_tf.clear()
            for b in bullets:
                p = body_tf.add_paragraph()
                p.text = b
                p.level = 0
                p.font.size = Pt(18)

        # ---- Images (placeholder safe – optional hook) ----
        # for ph in slide.placeholders:
        #     if ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
        #         generate_image_and_replace(ph, title, bullets)

    # --------------------------------------------
    # Re-append Thank You slide (unchanged)
    # --------------------------------------------
    prs.slides._sldIdLst.append(
        prs.slides._sldIdLst[-len(prs.slides)]
    )

    # --------------------------------------------
    # Save
    # --------------------------------------------
    os.makedirs("generated", exist_ok=True)
    out = f"generated/cognizant_{uuid.uuid4().hex[:6]}.pptx"
    prs.save(out)
    return out
