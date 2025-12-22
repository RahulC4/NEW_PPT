import os
import uuid
from pptx import Presentation
from pptx.util import Pt
from utils import text_client, image_client, get_env, logger

TEMPLATE_PATH = "templates/Cognizant.pptx"
IMAGE_MODEL = get_env("IMAGE_MODEL", required=True)


# --------------------------------------------------
# IMAGE GENERATION
# --------------------------------------------------
def generate_image(prompt: str) -> str:
    resp = image_client.images.generate(
        model=IMAGE_MODEL,
        prompt=prompt,
        size="1024x1024"
    )
    return resp.data[0].b64_json


# --------------------------------------------------
# FILL TEXT PLACEHOLDERS
# --------------------------------------------------
def fill_text(shape, text):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(18)


# --------------------------------------------------
# MAIN GENERATOR
# --------------------------------------------------
def generate_presentation(payload):
    slides = payload.get("slides") or payload.get("preview_slides", [])

    if not slides:
        raise ValueError("No slides provided")

    prs = Presentation(TEMPLATE_PATH)

    content_layout = prs.slide_layouts[1]  # Cognizant content layout

    for slide_data in slides:
        title = slide_data.get("title", "")
        bullets = slide_data.get("bullets", [])

        slide = prs.slides.add_slide(content_layout)

        # --------------------------------
        # TITLE
        # --------------------------------
        if slide.shapes.title:
            slide.shapes.title.text = title

        # --------------------------------
        # BODY BULLETS
        # --------------------------------
        body = slide.placeholders[1].text_frame
        body.clear()

        for b in bullets:
            p = body.add_paragraph()
            p.text = b
            p.level = 0
            p.font.size = Pt(18)

        # --------------------------------
        # IMAGE PLACEHOLDER (IF EXISTS)
        # --------------------------------
        for shape in slide.placeholders:
            if shape.placeholder_format.type == 18:  # PICTURE
                try:
                    img_b64 = generate_image(f"{title} business illustration")
                    img_path = f"/tmp/{uuid.uuid4().hex}.png"

                    with open(img_path, "wb") as f:
                        f.write(bytes.fromhex(img_b64))

                    slide.shapes.add_picture(
                        img_path,
                        shape.left,
                        shape.top,
                        shape.width,
                        shape.height
                    )
                except Exception:
                    logger.exception("Image generation failed")

    # --------------------------------
    # SAVE
    # --------------------------------
    os.makedirs("generated", exist_ok=True)
    out_path = f"generated/cognizant_{uuid.uuid4().hex[:6]}.pptx"
    prs.save(out_path)
    return out_path
