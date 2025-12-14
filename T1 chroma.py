import os
import uuid
import tempfile

from pptx import Presentation
from azure.storage.blob import BlobServiceClient
from chromadb import PersistentClient

from utils import get_env, logger, now_ts
from ppt.layout_extractor import extract_slide_layout_metadata
from search.search_utils import get_embedding

# ============================================================
# ENV CONFIG
# ============================================================
BLOB_CONN = get_env("AZURE_BLOB_CONN", required=True)
BLOB_CONTAINER = get_env("AZURE_BLOB_CONTAINER", "ppt-dataset")
CHROMA_PERSIST_DIR = get_env("CHROMA_PERSIST_DIR", "./chroma_db")

# ============================================================
# AZURE BLOB CLIENT
# ============================================================
blob_service = BlobServiceClient.from_connection_string(BLOB_CONN)
container_client = blob_service.get_container_client(BLOB_CONTAINER)

# ============================================================
# CHROMA CLIENT
# ============================================================
chroma_client = PersistentClient(path=CHROMA_PERSIST_DIR)

try:
    collection = chroma_client.get_collection("ppt_slides")
except Exception:
    collection = chroma_client.create_collection("ppt_slides")

# ============================================================
# HELPERS
# ============================================================
def extract_slide_text(slide):
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text and shape.text.strip():
            texts.append(shape.text.strip())
    return "\n".join(texts)


def ppt_already_indexed(ppt_name: str) -> bool:
    try:
        res = collection.query(where={"ppt_name": ppt_name}, n_results=1)
        return bool(res.get("ids") and res["ids"][0])
    except Exception:
        return False


# ============================================================
# MAIN INGESTION
# ============================================================
def process_blob(blob_name: str):
    logger.info(f"Processing PPT: {blob_name}")

    if ppt_already_indexed(blob_name):
        logger.info(f"Skipping {blob_name} (already indexed)")
        return

    # ---- Download PPT locally
    tmp_path = os.path.join(
        tempfile.gettempdir(),
        blob_name.replace("/", "_")
    )

    with open(tmp_path, "wb") as f:
        container_client.download_blob(blob_name).readinto(f)

    prs = Presentation(tmp_path)

    documents = []
    metadatas = []
    embeddings = []
    ids = []

    for idx, slide in enumerate(prs.slides):
        slide_text = extract_slide_text(slide)

        # ---- Extract layout metadata (SAFE)
        layout_meta = extract_slide_layout_metadata(prs, idx)

        metadata = {
            "ppt_name": blob_name,
            "slide_index": idx,
            "layout_name": layout_meta["layout_name"],
            "master_name": layout_meta["master_name"],
            "placeholders": layout_meta["placeholders"],
            "indexed_on": now_ts()
        }

        embedding = get_embedding(slide_text or layout_meta["layout_name"])

        if embedding is None:
            logger.warning(f"Skipping slide {idx} (embedding failed)")
            continue

        documents.append(slide_text)
        metadatas.append(metadata)
        embeddings.append(embedding)
        ids.append(str(uuid.uuid4()))

    if not documents:
        logger.warning(f"No valid slides found in {blob_name}")
        return

    collection.add(
        documents=documents,
        embeddings=embeddings,
        metadatas=metadatas,
        ids=ids
    )

    logger.info(f"Indexed {len(documents)} slides from {blob_name}")


# ============================================================
# ENTRY POINT
# ============================================================
def main():
    logger.info("Starting PPT ingestion into Chroma...")

    for blob in container_client.list_blobs():
        if blob.name.lower().endswith(".pptx"):
            try:
                process_blob(blob.name)
            except Exception as e:
                logger.exception(f"Failed processing {blob.name}: {e}")

    logger.info("Ingestion complete.")


if __name__ == "__main__":
    main()
