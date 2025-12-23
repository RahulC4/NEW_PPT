import os
import streamlit as st
from datetime import datetime
from pptx import Presentation
from utils import logger

st.set_page_config(page_title="5 - Generate PPT", layout="wide")
st.title("Step 5 — Generate & Download Presentation")

# ------------------------------------------------------------
# SESSION STATE SAFETY
# ------------------------------------------------------------
if "generated_ppts" not in st.session_state:
    st.session_state["generated_ppts"] = []

if "ppt_generated_once" not in st.session_state:
    st.session_state["ppt_generated_once"] = False

payload = st.session_state.get("generation_payload")
theme = st.session_state.get("ppt_theme", "auto")

if not payload:
    st.warning("No generation payload found. Please complete Preview first.")
    st.stop()

# ------------------------------------------------------------
# HELPER — Extract title from first slide
# ------------------------------------------------------------
def extract_title_from_ppt(ppt_path):
    try:
        prs = Presentation(ppt_path)
        if prs.slides and prs.slides[0].shapes.title:
            return prs.slides[0].shapes.title.text.strip()
    except Exception:
        pass
    return "Generated_Presentation"

# ------------------------------------------------------------
# GENERATE PPT (ONLY ONCE)
# ------------------------------------------------------------
if not st.session_state["ppt_generated_once"]:
    with st.spinner("Generating final PowerPoint..."):
        try:
            # ---------------- Generate PPT ----------------
            if theme == "cognizant":
                from generate_ppt_cognizant import generate_presentation_cognizant
                out_path = generate_presentation_cognizant(payload)
            else:
                from generate_ppt_llm import generate_presentation
                out_path = generate_presentation(payload)

            # ---------------- Rename logic ----------------
            ppt_title = extract_title_from_ppt(out_path)
            timestamp = datetime.now().strftime("%d_%b_%H-%M")
            display_name = f"{ppt_title}_{timestamp}.pptx"

            # ---------------- Store in session (LATEST FIRST) ----------------
            st.session_state["generated_ppts"].insert(
                0,
                {
                    "path": out_path,
                    "name": display_name,
                    "created_at": datetime.now(),
                }
            )

            # 🔒 Lock generation (prevents duplicate on download)
            st.session_state["ppt_generated_once"] = True

            st.success("✅ PPT generated successfully!")

        except Exception as e:
            logger.exception("PPT generation failed")
            st.error(f"Failed to generate PPT: {e}")
            st.stop()

# ------------------------------------------------------------
# DISPLAY GENERATED PPTs (SESSION ONLY)
# ------------------------------------------------------------
st.markdown("---")
st.subheader("📂 Generated PPTs (This Session)")

if not st.session_state["generated_ppts"]:
    st.caption("No PPTs generated yet.")
else:
    for idx, item in enumerate(st.session_state["generated_ppts"]):
        col1, col2 = st.columns([4, 2])

        with col1:
            st.write(f"**{idx + 1}. {item['name']}**")

        with col2:
            try:
                with open(item["path"], "rb") as f:
                    st.download_button(
                        label="⬇️ Download",
                        data=f,
                        file_name=item["name"],
                        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                        key=f"download_{idx}",
                    )
            except Exception:
                st.caption("File not available on disk.")

# ------------------------------------------------------------
# NAVIGATION
# ------------------------------------------------------------
st.markdown("---")

if st.button("⬅ Back to Home"):
    # Reset generation lock for next flow
    st.session_state["ppt_generated_once"] = False
    st.switch_page("pages/1_Home.py")
