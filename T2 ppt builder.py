import json
from pptx import Presentation
from utils import logger


def _find_layout_by_name(prs: Presentation, layout_name: str):
    for layout in prs.slide_layouts:
        if layout.name == layout_name:
            return layout
    return prs.slide_layouts[0]  # fallback


def _fill_placeholders(slide, placeholders_meta, content_map):
    """
    Fill slide placeholders using placeholder idx mapping.
    """
    for shape in slide.placeholders:
        idx = shape.placeholder_format.idx
        key = f"ph_{idx}"

        if key in content_map:
            try:
                shape.text = content_map[key]
            except Exception:
                pass


def build_presentation(mapped_templates, content_profile, output_path):
    """
    Build final PPT using layout_name + placeholder metadata.
    """

    logger.info("Building presentation...")
    prs = Presentation()

    for slide_plan in mapped_templates:
        layout_name = slide_plan["layout_name"]
        placeholders_meta = json.loads(slide_plan["placeholders"])

        layout = _find_layout_by_name(prs, layout_name)
        slide = prs.slides.add_slide(layout)

        # Build content map
        content_map = {}
        for ph in placeholders_meta:
            idx = ph["idx"]
            name = ph["name"].lower()

            # heuristic mapping
            if "title" in name:
                content_map[f"ph_{idx}"] = content_profile.get("title", "")
            elif "content" in name or "body" in name:
                content_map[f"ph_{idx}"] = content_profile.get(
                    slide_plan["section"], ""
                )

        _fill_placeholders(slide, placeholders_meta, content_map)

    prs.save(output_path)
    logger.info(f"PPT generated successfully: {output_path}")
