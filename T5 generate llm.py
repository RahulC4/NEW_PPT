import os
import streamlit as st
from pptx import Presentation
from pptx.util import Pt
import uuid

st.set_page_config(page_title="5 - Generate PPT", layout="wide")
st.title("5 — Generate Final PPT")

slides = st.session_state.get("final_slides")
if not slides:
    st.error("No preview content found")
    st.stop()

prs = Presentation()

for slide in slides:
    if slide["type"] == "title":
        s = prs.slides.add_slide(prs.slide_layouts[6])
        box = s.shapes.add_textbox(
            prs.slide_width * 0.15,
            prs.slide_height * 0.4,
            prs.slide_width * 0.7,
            Pt(80),
        )
        p = box.text_frame.paragraphs[0]
        p.text = slide["title"]
        p.font.size = Pt(40)
        p.font.bold = True
        p.alignment = 1
        continue

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = slide["title"]
    body = s.placeholders[1].text_frame
    body.clear()

    for b in slide["bullets"]:
        para = body.add_paragraph()
        para.text = b
        para.font.size = Pt(18)

# Thank You
end = prs.slides.add_slide(prs.slide_layouts[1])
end.shapes.title.text = "Thank You"

out_path = f"generated/ppt_{uuid.uuid4().hex[:6]}.pptx"
prs.save(out_path)

st.success("PPT Generated!")
with open(out_path, "rb") as f:
    st.download_button(
        "⬇️ Download PPT",
        f,
        file_name=os.path.basename(out_path),
        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
