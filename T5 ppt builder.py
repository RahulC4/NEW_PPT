from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Pt
from utils import logger


def build_presentation(mapped_templates, content_profile, output_path):
    """
    Builds a PowerPoint using mapped templates and generated content.
    """

    prs = Presentation()

    for idx, template in enumerate(mapped_templates):
        slide_layout = prs.slide_layouts[template["layout_id"]]
        slide = prs.slides.add_slide(slide_layout)

        slide_content = content_profile[idx]

        fill_placeholders(slide, slide_content)

    prs.save(output_path)
    logger.info(f"PPT saved to {output_path}")


def fill_placeholders(slide, content_map):
    """
    Populates slide placeholders with generated content.
    """

    for shape in slide.placeholders:
        if not shape.has_text_frame:
            continue

        ph_type = shape.placeholder_format.type
        tf = shape.text_frame
        tf.clear()

        # --------------------
        # TITLE PLACEHOLDER
        # --------------------
        if ph_type == PP_PLACEHOLDER.TITLE:
            title = content_map.get("title")
            if title:
                tf.text = title

        # --------------------
        # BODY / CONTENT
        # --------------------
        elif ph_type in (
            PP_PLACEHOLDER.BODY,
            PP_PLACEHOLDER.CONTENT,
            PP_PLACEHOLDER.OBJECT,
        ):
            bullets = content_map.get("bullets")

            if isinstance(bullets, list) and bullets:
                for i, bullet in enumerate(bullets):
                    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                    p.text = bullet
                    p.level = 0
                    for run in p.runs:
                        run.font.size = Pt(18)
            else:
                tf.text = ""

        # --------------------
        # SUBTITLE (OPTIONAL)
        # --------------------
        elif ph_type == PP_PLACEHOLDER.SUBTITLE:
            subtitle = content_map.get("subtitle")
            if subtitle:
                tf.text = subtitle
