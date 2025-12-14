from pptx import Presentation
from utils import logger


def extract_slide_layout_metadata(prs: Presentation, slide_index: int):
    """
    Extract stable, python-pptx-safe layout metadata.
    """

    slide = prs.slides[slide_index]
    layout = slide.slide_layout  # ✅ already SlideLayout

    layout_info = {
        "layout_name": layout.name,
        "master_name": layout.slide_master.name,
        "placeholders": []
    }

    for ph in layout.placeholders:
        layout_info["placeholders"].append({
            "idx": ph.placeholder_format.idx,
            "type": str(ph.placeholder_format.type),
            "name": ph.name
        })

    return layout_info
