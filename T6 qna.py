# generate_ppt.py
import os
import uuid
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt
from utils import logger


def replace_text_in_shape(shape, new_text):
    if not shape.has_text_frame:
        return

    tf = shape.text_frame
    tf.clear()

    for i, line in enumerate(new_text.split("\n")):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = line.strip()
        p.font.size = Pt(18)
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        p.bullet = True if len(new_text.split("\n")) > 1 else False

    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE


def clone_slide(prs, source_slide):
    new_slide = prs.slides.add_slide(prs.slide_layouts[6])
    for shape in source_slide.shapes:
        new_el = shape.element.clone()
        new_slide.shapes._spTree.insert_element_before(new_el, "p:extLst")
    return new_slide


def generate_presentation(payload):
    logger.info("Generating final presentation...")

    selected_slides = payload["selected_slides"]
    answers_map = payload["answers_map"]

    prs_out = Presentation()

    for slide_struct in selected_slides:
        ppt_path = slide_struct["ppt_path"]
        slide_index = slide_struct["slide_index"]
        editable_shapes = slide_struct["editable_shapes"]

        prs_src = Presentation(ppt_path)
        src_slide = prs_src.slides[slide_index]
        new_slide = clone_slide(prs_out, src_slide)

        slide_answers = answers_map.get(slide_struct["slide_id"], {}).get("answers", {})

        for shape_entry in editable_shapes:
            original_text = shape_entry["text"]

            for shp in new_slide.shapes:
                if shp.has_text_frame and shp.text.strip() == original_text:
                    replacement = "\n".join(
                        v for v in slide_answers.values() if v.strip()
                    )
                    if replacement:
                        replace_text_in_shape(shp, replacement)
                    break

    out_path = os.path.join(
        "generated",
        f"ppt_{uuid.uuid4().hex[:6]}.pptx"
    )

    prs_out.save(out_path)
    logger.info(f"PPT saved at {out_path}")
    return out_path
