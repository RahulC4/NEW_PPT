# ingestion/ingestion_chroma.py

import os
import uuid
import tempfile
from pptx import Presentation
from azure.storage.blob import BlobServiceClient
from openai import AzureOpenAI
from chromadb import PersistentClient

from utils import get_env, logger, now_ts
from llm.slide_classifier import classify_slide
from ppt.layout_extractor import extract_slide_layout_metadata

# ============================================================
# ENV + CONFIG
# ============================================================
BLOB_CONN = get_env("AZURE_BLOB_CONN", required=True)
BLOB_CONTAINER = get_env("AZURE_BLOB_CONTAINER", "ppt-dataset")

EMBEDDING_MODEL = get_env("EMBEDDING_MODEL", "text-embedding-3-large")
CHROMA_PERSIST_DIR = get_env("CHROMA_PERSIST_DIR", "./chroma_db")

# Azure OpenAI for embeddings
text_client = AzureOpenAI(
    azure_endpoint=get_env("OPENAI_API_BASE", required=True),
    api_key=get_env("OPENAI_API_KEY", required=True),
    api_version=get_env("OPENAI_API_VERSION", required=True)
)

# Azure Blob
blob_client = BlobServiceClient.from_connection_string(BLOB_CONN)
container_client = blob_client.get_container_client(BLOB_CONTAINER)

# Chroma DB
chroma_client = PersistentClient(path=CHROMA_PERSIST_DIR)
try:
    collection = chroma_client.get_collection("ppt_slides")
except:
    collection = chroma_client.create_collection("ppt_slides")

# ============================================================
# HELPERS
# ============================================================
def azure_embed_func(texts):
    """Generate embeddings using Azure OpenAI."""
    try:
        resp = text_client.embeddings.create(model=EMBEDDING_MODEL, input=texts)
        return [d.embedding for d in resp.data]
    except Exception as e:
        logger.exception(f"Embedding failed: {e}")
        return []


def extract_raw_text(slide):
    """Extract all visible text from a slide."""
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text and shape.text.strip():
            texts.append(shape.text.strip())
    return "\n".join(texts)


def ppt_already_indexed(ppt_name):
    """Check if slides for this PPT already exist in Chroma."""
    try:
        res = collection.query(where={"ppt_name": ppt_name}, n_results=1)
        ids = res.get("ids", [[]])[0]
        return len(ids) > 0
    except:
        return False


# ============================================================
# MAIN PROCESSOR
# ============================================================
def process_blob(blob_name):
    logger.info(f"Processing PPT: {blob_name}")

    # ---------------------------
    # 1. Download PPT file
    # ---------------------------
    tmp_path = os.path.join(tempfile.gettempdir(), blob_name.replace("/", "_"))
    with open(tmp_path, "wb") as fp:
        container_client.download_blob(blob_name).readinto(fp)

    prs = Presentation(tmp_path)

    # Skip if already indexed
    if ppt_already_indexed(blob_name):
        logger.info(f"Skipping {blob_name} — already indexed.")
        return

    docs = []
    metadatas = []
    ids = []

    # ---------------------------
    # 2. Extract text + layout + slide_type
    # ---------------------------
    for i, slide in enumerate(prs.slides):
        raw_text = extract_raw_text(slide)

        # ---- Slide Classification (AI) ----
        cls = classify_slide(raw_text)
        slide_type = cls["slide_type"]
        confidence = cls["confidence"]

        # ---- Layout Extraction ----
        layout_meta = extract_slide_layout_metadata(prs, i)

        slide_id = f"{os.path.splitext(os.path.basename(blob_name))[0]}_Slide_{i:02d}"

        metadata = {
            "ppt_name": blob_name,
            "slide_index": str(i),
            "slide_id": slide_id,
            "title": raw_text.split("\n", 1)[0] if raw_text else "",
            "raw_text": raw_text,
            "slide_type": slide_type,
            "slide_confidence": confidence,
            "layout": layout_meta,
            "indexed_on": now_ts()
        }

        docs.append(raw_text)
        metadatas.append(metadata)
        ids.append(str(uuid.uuid4()))

    # ---------------------------
    # 3. Embed slides
    # ---------------------------
    embeddings = azure_embed_func([m["raw_text"] for m in metadatas])

    if not embeddings or len(embeddings) != len(docs):
        logger.error(f"Embedding failed or mismatch for: {blob_name}")
        return

    # ---------------------------
    # 4. Store in Chroma
    # ---------------------------
    try:
        collection.add(
            documents=docs,
            embeddings=embeddings,
            metadatas=metadatas,
            ids=ids
        )
        logger.info(f"Indexed {len(docs)} slides from {blob_name}")
    except Exception as e:
        logger.exception(f"Failed to index slides from {blob_name}: {e}")


def main():
    logger.info("Starting PPT ingestion into Chroma...")
    for blob in container_client.list_blobs():
        if blob.name.lower().endswith(".pptx"):
            try:
                process_blob(blob.name)
            except Exception as e:
                logger.exception(f"Error processing {blob.name}: {e}")
    logger.info("Ingestion complete.")


if __name__ == "__main__":
    main()
