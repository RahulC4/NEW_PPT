# pages/5_Generate_PPT.py

import os
import streamlit as st
from datetime import datetime
from pptx import Presentation

from utils import logger
from azure_blob_utils import upload_ppt_to_blob

st.set_page_config(page_title="5 - Generate PPT", layout="wide")
st.title("Step 5 — Generate Your Presentation")

# ------------------------------------------------------------
# SAFE SESSION INIT
# ------------------------------------------------------------
st.session_state.setdefault("generated_ppts", [])

payload = st.session_state.get("generation_payload")
theme = st.session_state.get("ppt_theme", "auto")

if not payload:
    st.warning("No generation payload found. Complete Q&A first.")
    st.stop()

st.write("Generating final PPT...")

# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------
def extract_title_from_ppt(ppt_path):
    try:
        prs = Presentation(ppt_path)
        if prs.slides and prs.slides[0].shapes.title:
            return prs.slides[0].shapes.title.text.strip()
    except Exception:
        pass
    return "Generated_Presentation"


def safe_filename(name: str):
    return "".join(c for c in name if c.isalnum() or c in (" ", "_", "-")).strip()


# ------------------------------------------------------------
# GENERATE + UPLOAD
# ------------------------------------------------------------
try:
    # 1️⃣ Generate PPT locally
    if theme == "cognizant":
        from generate_ppt_cognizant import generate_presentation_cognizant
        local_ppt_path = generate_presentation_cognizant(payload)
    else:
        from generate_ppt_llm import generate_presentation
        local_ppt_path = generate_presentation(payload)

    # 2️⃣ Build final filename
    title = extract_title_from_ppt(local_ppt_path)
    title = safe_filename(title)
    timestamp = datetime.now().strftime("%d_%b_%Y_%H-%M")
    final_name = f"{title}_{timestamp}.pptx"

    # 3️⃣ Upload to Azure Blob (generated-ppts container)
    upload_ppt_to_blob(
        local_ppt_path,
        blob_name=final_name,
        container_name="generated-ppts"
    )

    # 4️⃣ Store in session (latest first)
    st.session_state["generated_ppts"].insert(
        0,
        {
            "blob_name": final_name,
            "created_at": datetime.now(),
            "theme": theme,
        }
    )

    st.success("✅ PPT generated and uploaded successfully!")
    st.write(f"**File:** `{final_name}`")

except Exception as e:
    logger.exception("Generation failed")
    st.error(f"Failed to generate PPT: {e}")
    st.stop()

# ------------------------------------------------------------
# SESSION GENERATED PPTs
# ------------------------------------------------------------
st.markdown("---")
st.subheader("📂 Generated PPTs (This Session)")

if not st.session_state["generated_ppts"]:
    st.caption("No PPTs generated yet.")
else:
    for i, item in enumerate(st.session_state["generated_ppts"], start=1):
        st.write(f"{i}. {item['blob_name']}  ·  ({item['theme']})")

# ------------------------------------------------------------
# NAVIGATION
# ------------------------------------------------------------
if st.button("Back to Home"):
    st.switch_page("pages/1_Home.py")
