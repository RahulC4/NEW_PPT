# ============================================================
# generate_ppt_cognizant.py
# Cognizant Template | Preview-driven | TEXT-ONLY
# ============================================================

import os
import uuid
from datetime import datetime
from copy import deepcopy

from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor

from utils import logger


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
COGNIZANT_TEMPLATE = os.path.join(TEMPLATES_DIR, "Cognizant.pptx")


# ------------------------------------------------------------
# SLIDE CLONE (SAFE)
# ------------------------------------------------------------
def clone_slide(prs, slide):
    new_slide = prs.slides.add_slide(slide.slide_layout)

    # Remove default shapes
    for shp in list(new_slide.shapes):
        new_slide.shapes._spTree.remove(shp._element)

    # Copy shapes from source slide
    for shp in slide.shapes:
        new_slide.shapes._spTree.insert_element_before(
            deepcopy(shp._element), "p:extLst"
        )

    return new_slide


# ------------------------------------------------------------
# FOOTER (MANUAL — RELIABLE)
# ------------------------------------------------------------
def add_footer(slide):
    year = datetime.now().year
    tb = slide.shapes.add_textbox(
        left=Inches(0.5),
        top=slide.slide_height - Inches(0.45),
        width=Inches(4),
        height=Pt(20),
    )

    p = tb.text_frame.paragraphs[0]
    p.text = f"© {year} Cognizant"
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(140, 140, 140)


# ------------------------------------------------------------
# TITLE SLIDE (UNCHANGED BEHAVIOR)
# ------------------------------------------------------------
def set_title_white_full_width(prs, slide, text):
    tf = slide.shapes.title.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    slide.shapes.title.left = Inches(0.5)
    slide.shapes.title.width = prs.slide_width - Inches(1)


# ------------------------------------------------------------
# CONTENT TITLE (FORCED WHITE)
# ------------------------------------------------------------
def set_content_title_white(slide, title):
    if not slide.shapes.title:
        return

    tf = slide.shapes.title.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)


# ------------------------------------------------------------
# CONTENT BULLETS (CUSTOM TEXTBOX – FIXES ALL ISSUES)
# ------------------------------------------------------------
def add_bullet_textbox(prs, slide, bullets):
    tb = slide.shapes.add_textbox(
        left=Inches(1),
        top=Inches(2),
        width=prs.slide_width - Inches(2),
        height=prs.slide_height - Inches(3),
    )

    tf = tb.text_frame
    tf.word_wrap = True
    tf.clear()

    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(20)


# ------------------------------------------------------------
# MAIN GENERATOR
# ------------------------------------------------------------
def generate_presentation_cognizant(payload):
    slides = payload.get("slides")
    if not slides:
        raise ValueError("No preview slides found")

    prs = Presentation(COGNIZANT_TEMPLATE)

    # Masters
    title_master = prs.slides[0]
    content_master = prs.slides[3]   # 👈 YOUR 4th SLIDE DESIGN
    thankyou_master = prs.slides[-1]

    # Remove all existing slides
    for i in reversed(range(len(prs.slides))):
        slide_id = prs.slides._sldIdLst[i].rId
        prs.part.drop_rel(slide_id)
        del prs.slides._sldIdLst[i]

    # --------------------------------------------------------
    # TITLE SLIDE
    # --------------------------------------------------------
    title_slide = clone_slide(prs, title_master)
    set_title_white_full_width(prs, title_slide, slides[0]["title"])
    add_footer(title_slide)

    # --------------------------------------------------------
    # CONTENT SLIDES (FIXED)
    # --------------------------------------------------------
    for slide_data in slides[1:]:
        slide = clone_slide(prs, content_master)

        title = slide_data.get("title", "")
        bullets = [
            b.strip()
            for b in (slide_data.get("bullets") or [])
            if isinstance(b, str) and b.strip()
        ]

        set_content_title_white(slide, title)
        add_bullet_textbox(prs, slide, bullets)
        add_footer(slide)

    # --------------------------------------------------------
    # THANK YOU SLIDE
    # --------------------------------------------------------
    thank_slide = clone_slide(prs, thankyou_master)
    if thank_slide.shapes.title:
        thank_slide.shapes.title.text = "Thank You"
    add_footer(thank_slide)

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------
    os.makedirs("generated", exist_ok=True)
    out = f"generated/cognizant_{uuid.uuid4().hex[:6]}.pptx"
    prs.save(out)
    return out
