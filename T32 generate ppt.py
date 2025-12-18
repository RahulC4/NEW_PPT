# ============================================================
# generate_ppt.py – FINAL FIXED VERSION
# ============================================================

import os
import uuid
import json
import tempfile
from typing import Dict, List, Any

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

from utils import get_env, logger, text_client

# ------------------------------------------------------------
CHAT_MODEL = get_env("CHAT_MODEL", required=True)
# ------------------------------------------------------------


# =========================
# Helpers
# =========================

def _answers_to_text(qna: Dict[str, str]) -> str:
    """Convert Q&A dict into readable text block for LLM"""
    return "\n".join(
        f"Q: {q}\nA: {a}"
        for q, a in qna.items()
        if a and str(a).strip()
    )


def _llm_generate_bullets(reference_text: str, qna_text: str) -> List[str]:
    """
    LLM generates ONLY bullets.
    Titles are forbidden here.
    """

    system_prompt = (
        "You are a senior presentation content writer.\n"
        "Generate 3–6 concise bullet points.\n"
        "Rules:\n"
        "- DO NOT generate a slide title\n"
        "- DO NOT copy reference text verbatim\n"
        "- Use Q&A as primary source\n"
        "- Reference content is context only\n"
        "- Skip unanswered questions\n"
    )

    user_prompt = (
        f"Reference context:\n{reference_text}\n\n"
        f"User Q&A:\n{qna_text}\n\n"
        "Generate bullet points only."
    )

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            temperature=0.6,
            max_completion_tokens=500,
        )

        raw = resp.choices[0].message.content.strip()
        bullets = [
            line.lstrip("-• ").strip()
            for line in raw.splitlines()
            if line.strip()
        ]
        return bullets[:6] or ["Key points derived from provided inputs."]

    except Exception as e:
        logger.warning(f"LLM bullet generation failed: {e}")
        return ["Key points derived from provided inputs."]


# =========================
# PPT Builders
# =========================

def _add_title_slide(prs: Presentation, title_text: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    box = slide.shapes.add_textbox(
        Inches(1.5),
        Inches(3),
        prs.slide_width - Inches(3),
        Inches(2),
    )

    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 102, 204)
    p.alignment = PP_ALIGN.CENTER


def _add_agenda_slide(prs: Presentation, agenda_items: List[str]):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Agenda"

    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()

    for item in agenda_items:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(20)
        p.level = 0


def _add_content_slide(
    prs: Presentation,
    title: str,
    bullets: List[str],
):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title

    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()

    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.font.size = Pt(18)
        p.level = 0


def _add_thank_you_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(
        Inches(2),
        Inches(3),
        prs.slide_width - Inches(4),
        Inches(2),
    )

    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "Thank You"
    p.font.size = Pt(32)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER


# =========================
# MAIN ENTRY
# =========================

def generate_presentation(payload: Dict[str, Any]) -> str:
    """
    FINAL LOGIC:
    - Slide 1  : Title (Q&A answer ONLY)
    - Slide 2  : Agenda (reference titles)
    - Slides 3+: Content (title from reference, bullets from LLM using Q&A)
    - Last     : Thank You
    """

    prs = Presentation()

    selected_slides = payload["selected_slide_structs"]
    qna_global = payload.get("qna_answers", {})
    answers_by_slide = payload.get("answers_by_slide", {})

    # -----------------------------
    # 1️⃣ TITLE SLIDE
    # -----------------------------
    presentation_title = (
        qna_global.get("What should be the title of this presentation?")
        or "Presentation"
    )

    _add_title_slide(prs, presentation_title)

    # -----------------------------
    # 2️⃣ AGENDA SLIDE
    # -----------------------------
    content_slides = [
        s for s in selected_slides
        if s.get("slide_type") == "content"
    ]

    agenda_titles = [s["title"] for s in content_slides]
    _add_agenda_slide(prs, agenda_titles)

    # -----------------------------
    # 3️⃣ CONTENT SLIDES
    # -----------------------------
    for slide in content_slides:
        slide_id = slide["slide_id"]

        # 🔒 STRICT: title ONLY from reference slide
        slide_title = slide["title"]

        qna_text = _answers_to_text(
            answers_by_slide.get(slide_id, {})
        )

        bullets = _llm_generate_bullets(
            reference_text=slide.get("text", ""),
            qna_text=qna_text,
        )

        _add_content_slide(
            prs,
            title=slide_title,
            bullets=bullets,
        )

    # -----------------------------
    # 4️⃣ THANK YOU
    # -----------------------------
    _add_thank_you_slide(prs)

    # -----------------------------
    # SAVE
    # -----------------------------
    out_path = os.path.join(
        tempfile.gettempdir(),
        f"generated_{uuid.uuid4().hex[:8]}.pptx",
    )
    prs.save(out_path)

    return out_path
