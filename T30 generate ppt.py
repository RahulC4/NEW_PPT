# ============================================================
# generate_ppt.py – Corporate path (FINAL)
# Title → Agenda → Content → Thank You
# Content generated primarily from QnA, reference is guidance only
# ============================================================

import os
import re
import uuid
import json
import tempfile
from typing import Dict, List, Any

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from utils import get_env, logger, now_ts, ensure_dir, text_client
from azure_blob_utils import upload_ppt_to_blob, upload_json_to_blob

# Optional Streamlit session fallback
try:
    import streamlit as st
    _SESSION = st.session_state
except Exception:
    _SESSION = {}

ensure_dir("generated")
CHAT_MODEL = get_env("CHAT_MODEL", required=True)

# ============================================================
# Helpers
# ============================================================

def _from_session(key, default=None):
    try:
        return _SESSION.get(key, default)
    except Exception:
        return default


def _normalize_list(x):
    if not x:
        return []
    if isinstance(x, list):
        return x
    if isinstance(x, (tuple, set)):
        return list(x)
    return [x]


def _resolve_selection(payload: Dict[str, Any]) -> List[Dict[str, Any]]:
    if payload.get("selected_slide_structs"):
        return payload["selected_slide_structs"]

    if _from_session("selected_slide_structs"):
        return _from_session("selected_slide_structs")

    raise ValueError("No selected slides found")


def _extract_answers_map(payload: Dict[str, Any]) -> Dict[str, Dict[str, str]]:
    return (
        payload.get("answers_map")
        or payload.get("answers_by_slide")
        or _from_session("answers_by_slide")
        or {}
    )


def _compose_guidance_text(slide: Dict[str, Any]) -> str:
    return (slide.get("text") or slide.get("title") or "")[:1200]


def _clean_answers_block(answers: Dict[str, str]) -> str:
    if not isinstance(answers, dict):
        return ""

    cleaned = []
    for q, a in answers.items():
        if not a:
            continue
        if isinstance(a, str) and a.strip().lower() in {"not sure", "na", "n/a"}:
            continue
        cleaned.append(f"{q}: {a}")

    return "\n".join(cleaned)


# ============================================================
# LLM – Content Slide Generator
# ============================================================

def _llm_generate_content_slide(
    guidance_text: str,
    per_slide_answers: Dict[str, str],
    title_hint: str,
) -> Dict[str, Any]:

    answers_block = _clean_answers_block(per_slide_answers)

    system_prompt = f"""
You are a senior presentation consultant.

TASK:
Create ONE professional PowerPoint content slide.

RULES:
- Use USER Q&A as the PRIMARY source
- Use reference text only for context
- Ignore unanswered or unclear questions
- Do NOT copy reference text
- Produce 3–6 concise business bullets
- Bullets do NOT need to map 1:1 to questions

REFERENCE (guidance only):
{guidance_text}

USER Q&A (primary):
{answers_block}

FORMAT:
Title: <slide title>
- Bullet
- Bullet
- Bullet
"""

    try:
        resp = text_client.chat.completions.create(
            model=CHAT_MODEL,
            messages=[{"role": "system", "content": system_prompt}],
            temperature=0.9,
            max_completion_tokens=700,
        )
        raw = resp.choices[0].message.content.strip()
    except Exception as e:
        logger.warning(f"LLM failed, fallback used: {e}")
        return {
            "title": title_hint,
            "bullets": [
                "Key objectives and scope defined.",
                "Approach aligned with business priorities.",
                "Next steps identified for execution.",
            ],
        }

    lines = [l.strip() for l in raw.splitlines() if l.strip()]
    title = title_hint
    bullets = []

    for line in lines:
        if line.lower().startswith("title"):
            title = line.split(":", 1)[-1].strip()
        elif line.startswith("-"):
            bullets.append(line[1:].strip())

    if not bullets:
        bullets = [
            "Context and objectives summarized.",
            "Solution approach synthesized.",
            "Actionable direction identified.",
        ]

    return {
        "title": title,
        "bullets": bullets[:6],
    }


# ============================================================
# PPT Builder
# ============================================================

def _build_corporate_presentation(
    selected_slides: List[Dict[str, Any]],
    answers_map: Dict[str, Dict[str, str]],
) -> str:

    prs = Presentation()

    # --------------------------------------------------------
    # 1️⃣ TITLE SLIDE (from QnA)
    # --------------------------------------------------------
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])

    title_answer = answers_map.get(
        selected_slides[0]["slide_id"], {}
    ).get("What should be the title of this presentation?", "Presentation")

    title_slide.shapes.title.text = title_answer
    title_slide.placeholders[1].text = "Auto-generated using AI"

    # --------------------------------------------------------
    # 2️⃣ GENERATE CONTENT SLIDES (LLM)
    # --------------------------------------------------------
    content_slides = []

    for slide in selected_slides:
        slide_type = slide.get("slide_type", "content")
        if slide_type in {"title", "agenda", "thankyou"}:
            continue

        slide_id = slide["slide_id"]

        content = _llm_generate_content_slide(
            guidance_text=_compose_guidance_text(slide),
            per_slide_answers=answers_map.get(slide_id, {}),
            title_hint=slide.get("title") or "Overview",
        )
        content_slides.append(content)

    # --------------------------------------------------------
    # 3️⃣ AGENDA SLIDE
    # --------------------------------------------------------
    agenda_slide = prs.slides.add_slide(prs.slide_layouts[1])
    agenda_slide.shapes.title.text = "Agenda"

    agenda_tf = agenda_slide.placeholders[1].text_frame
    agenda_tf.clear()

    for c in content_slides:
        p = agenda_tf.add_paragraph()
        p.text = c["title"]
        p.level = 0

    # --------------------------------------------------------
    # 4️⃣ CONTENT SLIDES
    # --------------------------------------------------------
    for c in content_slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = c["title"]

        tf = s.placeholders[1].text_frame
        tf.clear()

        for b in c["bullets"]:
            p = tf.add_paragraph()
            p.text = b
            p.level = 0

    # --------------------------------------------------------
    # 5️⃣ THANK YOU SLIDE
    # --------------------------------------------------------
    thank_you = prs.slides.add_slide(prs.slide_layouts[0])
    thank_you.shapes.title.text = "Thank You"
    thank_you.placeholders[1].text = "Questions?"

    out_path = os.path.join(
        tempfile.gettempdir(),
        f"generated_{uuid.uuid4().hex[:8]}.pptx"
    )
    prs.save(out_path)
    return out_path


# ============================================================
# PUBLIC API
# ============================================================

def generate_presentation(payload: Dict[str, Any]):
    selected_slides = _resolve_selection(payload)
    answers_map = _extract_answers_map(payload)

    out_path = _build_corporate_presentation(
        selected_slides=selected_slides,
        answers_map=answers_map,
    )

    # Optional upload + log
    try:
        fname = os.path.basename(out_path)
        with open(out_path, "rb") as f:
            upload_ppt_to_blob(f.read(), fname)

        upload_json_to_blob(
            json.dumps(
                {
                    "timestamp": now_ts(),
                    "slides": len(selected_slides),
                    "ppt": fname,
                    "error": False,
                },
                indent=2,
            ).encode("utf-8"),
            f"logs/{fname}.json",
        )
    except Exception as e:
        logger.warning(f"Upload skipped: {e}")

    return out_path
