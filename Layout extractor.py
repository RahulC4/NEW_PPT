# ppt/layout_extractor.py

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from utils import logger


def extract_slide_layout_metadata(prs: Presentation, slide_index: int):
    """
    Extracts font, placeholder, shape, and layout structure from a slide.
    Returns a JSON-serializable dict.
    """

    try:
        slide = prs.slides[slide_index]
    except:
        logger.error(f"Slide index {slide_index} out of range.")
        return {}

    layout_info = {
        "layout_id": slide.slide_layout.slide_layout_id,
        "master_id": slide.slide_layout.master.id,
        "placeholders": []
    }

    # ------------------------------
    # Extract Shape / Placeholder Info
    # ------------------------------
    for shape in slide.shapes:

        shape_info = {
            "is_placeholder": shape.is_placeholder,
            "placeholder_type": None,
            "shape_type": None,
            "text": "",
            "left": int(shape.left),
            "top": int(shape.top),
            "width": int(shape.width),
            "height": int(shape.height),
            "font": {},
            "has_text_frame": hasattr(shape, "text_frame"),
        }

        # Placeholder type
        if shape.is_placeholder:
            ph = shape.placeholder_format
            shape_info["placeholder_type"] = str(ph.type)

        # Shape type
        try:
            shape_info["shape_type"] = shape.shape_type
        except:
            shape_info["shape_type"] = None

        # Extract text
        if hasattr(shape, "text_frame") and shape.text_frame:
            text = shape.text_frame.text or ""
            shape_info["text"] = text.strip()

            # Extract first paragraph font settings
            try:
                p = shape.text_frame.paragraphs[0]
                if p.runs and len(p.runs) > 0:
                    r = p.runs[0]

                    shape_info["font"] = {
                        "name": r.font.name,
                        "size": r.font.size.pt if r.font.size else None,
                        "bold": r.font.bold,
                        "italic": r.font.italic,
                        "color": str(r.font.color.rgb) if r.font.color and r.font.color.rgb else None,
                    }
            except Exception as e:
                logger.warning(f"Font extraction failed: {e}")

        layout_info["placeholders"].append(shape_info)

    return layout_info



def extract_layouts_from_ppt(local_path: str):
    """
    Extracts layout metadata for all slides in a PPT.
    Returns list of slide layout dicts.
    For ingestion_chroma to store in Chroma metadata.
    """

    try:
        prs = Presentation(local_path)
    except Exception as e:
        logger.error(f"Failed to open PPT for layout extraction: {e}")
        return []

    all_slides = []

    for i in range(len(prs.slides)):
        layout_meta = extract_slide_layout_metadata(prs, i)
        all_slides.append(layout_meta)

    return all_slides
